"""Excel (.xlsx) export — M40."""
import io

import pytest

from backend.engine import run_model
from backend.export_xlsx import run_to_xlsx
from conftest import load_model


def test_xlsx_steady_export():
    pytest.importorskip("openpyxl")
    r = run_model(load_model("cstr.gPROMS", "examples"))
    data = run_to_xlsx(r["results"], r["timeSeries"], r["messages"])
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["Results"]
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    assert rows[0] == ["variable", "value"]
    vals = dict((rows[i][0], rows[i][1]) for i in range(1, len(rows)))
    assert abs(vals["ca"] - 1.0 / 3.0) < 1e-9
    # messages sheet exists with content
    ws2 = wb["Messages"]
    assert ws2["A1"].value == "level"


def test_xlsx_dynamic_export():
    pytest.importorskip("openpyxl")
    r = run_model(load_model("tank.gPROMS", "examples"))
    data = run_to_xlsx(r["results"], r["timeSeries"], r["messages"])
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["Results"]
    header = [c.value for c in next(ws.iter_rows())]
    assert header[0] == "t"
    assert "h" in header
    # data rows: every time point is present (plus the final-value block)
    n_data = sum(1 for _ in ws.iter_rows()) - 1
    assert n_data >= len(r["timeSeries"]["t"])
